"""Synthetic workbook fixtures exercise the loader; no sample is user data."""

import json
from copy import deepcopy
from datetime import datetime

import pytest
from openpyxl import Workbook

from backend.calculations.excel_loader import load_workbooks
from backend.calculations.validation import validate_dataset


def book(tmp_path, sheets, filename="facts.xlsx"):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    path = tmp_path / filename
    workbook.save(path)
    workbook.close()
    return path


def codes(result):
    return {item["code"] for item in result["report"]["summary"]["issues"]}


def test_merges_russian_sheets_with_wide_dates_and_long_incoming(tmp_path):
    path = book(
        tmp_path,
        {
            "Остатки": [
                ["Артикул", "Наименование", "Поставщик", "Ед. изм.", "Остаток"],
                ["000123", "Кабель", "ИЭК", "м", "1 200,5"],
            ],
            "Закупки": [
                ["sku", "MOQ", "Кратность заказа", "Срок поставки"],
                ["000123", 5, 10, 7],
            ],
            "Продажи": [
                ["Артикул", datetime(2025, 1, 1), "02.2025", "март 2025"],
                ["000123", 31, 28, None],
            ],
            "Поступления": [
                ["sku", "Дата", "Количество"],
                ["000123", datetime(2026, 10, 2), 20],
                ["000123", "03.10.2026", 30],
            ],
            "Сезонность": [["Артикул", 1, 2, 3], ["000123", 1.1, 0.9, 1]],
        },
    )
    result = load_workbooks([path], [{"filename": "Пользователь.xlsx"}])
    assert result["report"]["valid"], result["report"]
    product = result["payload"]["products"][0]
    assert product["sku"] == "000123"
    assert product["stock"] == 1200.5
    assert product["moq"] == 5 and product["order_multiple"] == 10
    assert product["sales"] == [
        {"month": "2025-01-01", "quantity": 31, "stockout_days": None},
        {"month": "2025-02-01", "quantity": 28, "stockout_days": None},
        {"month": "2025-03-01", "quantity": None, "stockout_days": None},
    ]
    assert product["incoming_known"] is True
    assert product["incoming"] == [
        {"date": "2026-10-02", "quantity": 20},
        {"date": "2026-10-03", "quantity": 30},
    ]
    assert product["seasonality"] == {"1": 1.1, "2": 0.9, "3": 1}
    assert "missing_sales_quantity" in codes(result)
    json.dumps(result, allow_nan=False)


def test_long_sales_zero_is_observed_missing_is_not_zero(tmp_path):
    path = book(
        tmp_path,
        {
            "Sales": [
                ["SKU", "Name", "Supplier", "Month", "Quantity", "Stockout days"],
                ["01", "Item", "Vendor", "2025-01", 0, 0],
                ["01", "Item", "Vendor", "2025-02", None, 28],
            ]
        },
    )
    result = load_workbooks([path])
    assert result["report"]["valid"]
    product = result["payload"]["products"][0]
    assert product["sales"][0]["quantity"] == 0
    assert product["sales"][1]["quantity"] is None
    assert product["stock"] is None
    assert product["incoming_known"] is False
    assert {"missing_optional_field", "missing_incoming", "missing_unit"} <= codes(result)


def test_explicit_incoming_zero_is_known_but_positive_undated_is_unknown(tmp_path):
    path = book(
        tmp_path,
        {
            "Products": [
                ["sku", "name", "supplier", "incoming"],
                ["zero", "Item", "Vendor", 0],
                ["unknown", "Item", "Vendor", None],
                ["undated", "Item", "Vendor", 10],
            ]
        },
    )
    result = load_workbooks([path])
    assert result["report"]["valid"]
    products = {product["sku"]: product for product in result["payload"]["products"]}
    assert products["zero"]["incoming_known"] is True
    assert products["unknown"]["incoming_known"] is False
    assert products["undated"]["incoming_known"] is False


@pytest.mark.parametrize("value", ["NaN", "Infinity", "1e9999", "oops", -1, True])
def test_invalid_numbers_fail_without_nonfinite_json(tmp_path, value):
    path = book(
        tmp_path,
        {
            "Products": [
                ["sku", "name", "supplier", "stock"],
                ["001", "Item", "Vendor", value],
            ]
        },
    )
    result = load_workbooks([path])
    assert not result["report"]["valid"]
    assert codes(result) & {"invalid_number", "invalid_range"}
    assert not validate_dataset(result["payload"])["valid"]
    json.dumps(result, allow_nan=False)
    issue = next(
        item for item in result["report"]["summary"]["issues"] if item["severity"] == "error"
    )
    assert issue["source"] == "facts.xlsx" and issue["row"] == 2 and issue["sku"] == "001"


def test_duplicate_static_sku_and_conflicting_supplier_fail(tmp_path):
    path = book(
        tmp_path,
        {
            "Stock": [
                ["sku", "name", "supplier", "stock"],
                ["01", "Item", "A", 1],
                ["01", "Item", "B", 1],
            ]
        },
    )
    result = load_workbooks([path])
    assert not result["report"]["valid"]
    assert {"duplicate_sku", "conflicting_value"} <= codes(result)


def test_same_month_in_two_files_is_error(tmp_path):
    sheets = {
        "Sales": [
            ["sku", "name", "supplier", "month", "sales"],
            ["01", "Item", "A", "2025-01", 5],
        ]
    }
    paths = [book(tmp_path, sheets, "first.xlsx"), book(tmp_path, sheets, "second.xlsx")]
    result = load_workbooks(paths)
    assert not result["report"]["valid"]
    assert "duplicate_sales" in codes(result)


def test_missing_critical_fields_and_invalid_dates(tmp_path):
    path = book(
        tmp_path,
        {
            "Sales": [
                ["sku", "month", "sales"],
                ["01", "2025-15", 5],
            ]
        },
    )
    result = load_workbooks([path])
    assert not result["report"]["valid"]
    assert {"missing_required_field", "invalid_date"} <= codes(result)


def test_custom_alias_header_preamble_and_formula_cache(tmp_path):
    path = book(
        tmp_path,
        {
            "Products": [
                ["Export generated by ERP"],
                [],
                ["SKU", "Name", "Supplier", "Доступно, шт"],
                ["001", "Item", "Vendor", "=2+3"],
            ]
        },
    )
    result = load_workbooks([path], aliases={"stock": ["Доступно, шт"]})
    assert result["report"]["valid"]
    assert result["payload"]["products"][0]["stock"] is None
    assert "formula_without_cache" in codes(result)


def test_formatted_numeric_sku_preserves_zero_mask(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["sku", "name", "supplier", "stock"])
    sheet.append([123, "Item", "Vendor", 1])
    sheet["A2"].number_format = "000000"
    path = tmp_path / "formatted.xlsx"
    workbook.save(path)
    workbook.close()
    result = load_workbooks([path])
    assert result["payload"]["products"][0]["sku"] == "000123"


def test_delivery_date_header_disambiguates_generic_quantity(tmp_path):
    path = book(
        tmp_path,
        {
            "Sheet1": [
                ["sku", "name", "supplier", "Delivery date", "Quantity"],
                ["01", "Item", "Vendor", "2026-09-25", 4],
            ]
        },
    )
    result = load_workbooks([path])
    assert result["report"]["valid"]
    product = result["payload"]["products"][0]
    assert product["sales"] == []
    assert product["incoming"] == [{"date": "2026-09-25", "quantity": 4}]


def test_invalid_stockout_and_seasonality_are_rejected(tmp_path):
    path = book(
        tmp_path,
        {
            "Sales": [
                ["sku", "name", "supplier", "month", "sales", "stockout_days"],
                ["01", "Item", "Vendor", "2025-02", 4, 30],
            ],
            "Seasonality": [["sku", "month", "coefficient"], ["01", 2, 0]],
        },
    )
    result = load_workbooks([path])
    assert not result["report"]["valid"]
    assert {"invalid_stockout_days", "invalid_range"} <= codes(result)


def test_validate_dataset_checks_direct_payload_and_does_not_mutate(tmp_path):
    path = book(
        tmp_path,
        {
            "Products": [
                ["sku", "name", "supplier", "stock"],
                ["01", "Item", "Vendor", 2],
            ]
        },
    )
    payload = load_workbooks([path])["payload"]
    payload["products"][0]["stock"] = float("nan")
    payload["products"][0]["incoming_known"] = True
    payload["products"][0]["incoming"] = [{"date": None, "quantity": 5}]
    before = deepcopy(payload)
    report = validate_dataset(payload)
    assert not report["valid"]
    assert {"invalid_number", "inconsistent_incoming"} <= {
        item["code"] for item in report["summary"]["issues"]
    }
    assert payload["products"][0]["incoming"] == before["products"][0]["incoming"]
    json.dumps(report, allow_nan=False)


def test_no_files_no_headers_and_corrupt_file_have_reports(tmp_path):
    assert not load_workbooks([])["report"]["valid"]
    path = book(tmp_path, {"Sheet1": [["Unknown", "Thing"], ["data", 1]]})
    assert "missing_sku_header" in codes(load_workbooks([path]))
    corrupt = tmp_path / "broken.xlsx"
    corrupt.write_text("not an Excel file")
    assert "workbook_read_error" in codes(load_workbooks([corrupt]))


def test_manufacturer_is_not_silently_used_as_supplier(tmp_path):
    path = book(
        tmp_path,
        {
            "Products": [
                ["sku", "name", "Производитель", "stock"],
                ["01", "Item", "Factory", 2],
            ]
        },
    )
    result = load_workbooks([path])
    assert not result["report"]["valid"]
    assert result["payload"]["products"][0]["supplier"] is None


def test_sales_contradicting_full_month_stockout_are_invalid(tmp_path):
    path = book(
        tmp_path,
        {
            "Sales": [
                ["sku", "name", "supplier", "month", "sales", "stockout_days"],
                ["01", "Item", "Vendor", "2025-02", 4, 28],
            ]
        },
    )
    result = load_workbooks([path])
    assert not result["report"]["valid"]
    assert "contradictory_stockout" in codes(result)


def test_identical_incoming_requires_disambiguation_before_counting_twice(tmp_path):
    path = book(
        tmp_path,
        {
            "Incoming": [
                ["sku", "name", "supplier", "date", "quantity"],
                ["01", "Item", "Vendor", "2026-10-01", 10],
                ["01", "Item", "Vendor", "2026-10-01", 10],
            ]
        },
    )
    result = load_workbooks([path])
    assert not result["report"]["valid"]
    assert "ambiguous_duplicate_incoming" in codes(result)


@pytest.mark.parametrize(
    ("sheet", "extra_headers", "extra_values"),
    [
        ("Products", [], []),
        ("Остатки", ["stock", "date"], [5, "2026-09-15"]),
    ],
)
def test_duplicate_master_rows_fail_even_without_numeric_fields_or_with_date(
    tmp_path, sheet, extra_headers, extra_values
):
    row = ["A", "Товар", "ИЭК", *extra_values]
    path = book(
        tmp_path,
        {sheet: [["sku", "name", "supplier", *extra_headers], row, row]},
    )
    result = load_workbooks([path])
    assert not result["report"]["valid"]
    assert "duplicate_sku" in codes(result)


@pytest.mark.parametrize(
    ("sheet", "extra_headers", "extra_values", "conflicting_field"),
    [
        (
            "В пути",
            ["incoming_quantity", "quantity", "date"],
            [5, 500, "2026-09-15"],
            "incoming_quantity",
        ),
        (
            "В пути",
            ["incoming_quantity", "incoming_date", "date"],
            [5, "2026-09-15", "2026-10-15"],
            "incoming_date",
        ),
        (
            "Sales",
            ["sales", "quantity", "month"],
            [5, 500, "2026-09"],
            "sales_quantity",
        ),
        (
            "Sales",
            ["sales", "month", "date"],
            [5, "2026-09", "2026-10-15"],
            "month",
        ),
        ("Остатки", ["stock", "quantity"], [5, 500], "stock"),
    ],
)
def test_generic_column_remapping_reports_conflicting_explicit_column(
    tmp_path, sheet, extra_headers, extra_values, conflicting_field
):
    path = book(
        tmp_path,
        {
            sheet: [
                ["sku", "name", "supplier", *extra_headers],
                ["A", "Товар", "ИЭК", *extra_values],
            ]
        },
    )
    result = load_workbooks([path])
    assert not result["report"]["valid"]
    assert not validate_dataset(result["payload"])["valid"]
    matching = [
        item
        for item in result["report"]["summary"]["issues"]
        if item["code"] == "duplicate_column" and item["field"] == conflicting_field
    ]
    assert len(matching) == 1
    assert matching[0]["row"] == 1 and matching[0]["severity"] == "error"
