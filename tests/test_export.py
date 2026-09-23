from datetime import date, datetime, timezone
from io import BytesIO
from uuid import uuid4

from conftest import parameters
from openpyxl import load_workbook

from backend.models import Calculation, CalculationParameters, Recommendation
from backend.services.export import export_calculation


def test_formula_text_unknowns_and_sheet_names():
    rows = []
    for supplier in ["Параметры", "a/b", "a?b", "A?B", "x" * 60]:
        rows.append(
            Recommendation(
                sku="00001",
                name='=HYPERLINK("https://example.com")',
                supplier=supplier,
                unit=None,
                stock=None,
                incoming_in_period=None,
                forecast_demand=None,
                safety_stock=None,
                moq=None,
                order_multiple=None,
                recommended_qty=None,
                urgency="unknown",
                reason="+SUM(A1:A3)",
                warnings=["@warning", "-unknown"],
            )
        )
    result = Calculation(
        calculation_id=uuid4(),
        dataset_id=uuid4(),
        created_at=datetime.now(timezone.utc),
        demo=False,
        provider="test",
        parameters=CalculationParameters(**parameters()),
        sources=[],
        recommendations=rows,
        forecast_start=date(2026, 9, 23),
        forecast_end=date(2026, 10, 14),
        warnings=["=1+1"],
    )
    workbook = load_workbook(BytesIO(export_calculation(result)), data_only=False)
    assert len(workbook.sheetnames) == 6
    assert len({name.lower() for name in workbook.sheetnames}) == 6
    for sheet in list(workbook)[1:]:
        assert len(sheet.title) <= 31
        assert sheet["A2"].value == "00001"
        assert sheet["B2"].value.startswith("=HYPERLINK")
        assert sheet["B2"].data_type == "s"
        assert sheet["L2"].data_type == "s"
        assert sheet["D2"].value is None
        assert sheet["J2"].value is None
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
