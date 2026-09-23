import json
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from conftest import calculate, parameters, upload, workbook_bytes
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.main import create_app
from backend.services.calculations import CalculationAdapter


def test_health_and_openapi(client):
    health = client.get("/api/health").json()
    assert health["calculations"]["available"]
    assert health["ai"]["status"] == "disabled"
    schema = client.get("/openapi.json").json()
    assert set(schema["paths"]) == {
        "/api/health",
        "/api/upload",
        "/api/calculate",
        "/api/recommendations",
        "/api/chat",
        "/api/export",
    }
    assert client.get("/docs").status_code == 200


@pytest.mark.parametrize(
    "name,data",
    [
        ("data.csv", b"a,b"),
        ("fake.xlsx", b"not a workbook"),
        ("data.xlsx", b""),
    ],
)
def test_invalid_upload(client, name, data):
    response = client.post("/api/upload", files={"files": (name, data)})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_file"
    assert not list(client.app.state.storage.root.glob("uploads/*"))


def test_upload_limits(settings):
    settings.max_file_bytes = 100
    settings.max_upload_files = 1
    with TestClient(create_app(settings)) as client:
        response = client.post("/api/upload", files={"files": ("data.xlsx", workbook_bytes())})
        assert response.status_code == 413
        response = client.post(
            "/api/upload", files=[("files", ("a.xlsx", b"x")), ("files", ("b.xlsx", b"x"))]
        )
        assert response.status_code == 400
        assert "error" in response.json()
    settings.max_request_bytes = 100
    with TestClient(create_app(settings)) as client:
        response = client.post("/api/upload", content=b"a" * 101)
        assert response.status_code == 413
        assert response.json()["error"]["code"] == "request_too_large"


def test_uncompressed_limit(settings):
    settings.max_uncompressed_bytes = 1000
    with TestClient(create_app(settings)) as client:
        response = client.post("/api/upload", files={"files": ("data.xlsx", workbook_bytes())})
        assert response.status_code == 413


def test_actual_cell_limit(settings):
    settings.max_workbook_cells = 2
    with TestClient(create_app(settings)) as client:
        response = client.post("/api/upload", files={"files": ("data.xlsx", workbook_bytes())})
        assert response.status_code == 413


def test_fake_zip_rejected(client):
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as archive:
        archive.writestr("test.txt", "hello")
    response = client.post("/api/upload", files={"files": ("fake.xlsx", buffer.getvalue())})
    assert response.status_code == 422


def test_multiple_files_safe_names(client):
    response = client.post(
        "/api/upload",
        files=[
            ("files", ("../../sales.xlsx", workbook_bytes())),
            ("files", ("C:\\private\\sales.xlsx", workbook_bytes())),
        ],
    )
    assert response.status_code == 201
    body = response.json()
    assert "payload" not in body and "provider" not in body
    assert [source["filename"] for source in body["sources"]] == ["sales.xlsx", "sales.xlsx"]
    assert len({s["file_id"] for s in body["sources"]}) == 2
    assert body["validation"]["business_validation"] == "not_performed"
    paths = list(client.app.state.storage.root.glob("uploads/*/*.xlsx"))
    assert len(paths) == 2
    assert all(p.name != "sales.xlsx" for p in paths)


def test_missing_provider_no_demo_fallback(settings):
    settings.demo_mode = False
    settings.calculation_module = "backend.calculations.does_not_exist"
    with TestClient(create_app(settings)) as client:
        assert not client.get("/api/health").json()["calculations"]["available"]
        response = client.post("/api/upload", files={"files": ("data.xlsx", workbook_bytes())})
        assert response.status_code == 503
        assert response.json()["error"]["code"] == "calculation_unavailable"
        assert not list(settings.data_dir.glob("uploads/*"))


@pytest.mark.parametrize(
    "changes",
    [
        {"lead_time_days": -1},
        {"review_period_days": 0},
        {"safety_stock_days": "20"},
        {"lead_time_days": True},
        {"demand_growth_adjustment": 6},
        {"calculation_date": "yesterday"},
        {"unexpected": 1},
    ],
)
def test_invalid_parameters(client, changes):
    response = client.post(
        "/api/calculate", json={"dataset_id": str(uuid4()), "parameters": parameters(**changes)}
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "validation_error"


def test_isolation_filters_export_and_no_key(client):
    first, second = upload(client, "first"), upload(client, "second")
    a = calculate(client, first["dataset_id"], safety_stock_days=10)
    b = calculate(client, second["dataset_id"], supplier="ИЭК", safety_stock_days=20)
    assert a["calculation_id"] != b["calculation_id"]
    assert a["sources"] == first["sources"]
    assert b["parameters"]["safety_stock_days"] == 20
    assert a["recommendations"][0]["recommended_qty"] != b["recommendations"][0]["recommended_qty"]
    response = client.get(
        "/api/recommendations",
        params={
            "calculation_id": a["calculation_id"],
            "supplier": "SystemElectric",
            "urgency": "high",
        },
    )
    assert response.json()["total"] == 1
    assert response.json()["dataset_id"] == first["dataset_id"]
    chat = client.post(
        "/api/chat",
        json={
            "dataset_id": first["dataset_id"],
            "calculation_id": a["calculation_id"],
            "message": "Почему?",
        },
    )
    assert chat.status_code == 200
    assert chat.json()["ai_available"] is False
    assert "ИИ не подключён" in chat.json()["message"]
    assert "ДЕМО" in chat.json()["message"]
    export = client.get(chat.json()["export_url"])
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content))
    assert set(workbook.sheetnames) == {"Параметры", "ИЭК", "SystemElectric"}
    assert workbook["ИЭК"]["A2"].value == "000123"
    assert workbook["ИЭК"]["A2"].data_type == "s"
    assert workbook["ИЭК"]["J2"].value == a["recommendations"][0]["recommended_qty"]
    assert workbook["ИЭК"].freeze_panes == "A2"
    assert workbook["ИЭК"].auto_filter.ref == "A1:M2"
    metadata = dict(workbook["Параметры"].values)
    assert metadata["calculation_id"] == a["calculation_id"]
    assert metadata["safety_stock_days"] == 10
    assert "ДЕМО" in metadata["Режим"]
    workbook.close()


def test_chat_dataset_and_conversation_scope(client):
    a, b = upload(client), upload(client)
    calculation = calculate(client, a["dataset_id"])
    response = client.post(
        "/api/chat",
        json={
            "dataset_id": b["dataset_id"],
            "calculation_id": calculation["calculation_id"],
            "message": "Объясни",
        },
    )
    assert response.status_code == 409
    conversation = client.post(
        "/api/chat", json={"dataset_id": a["dataset_id"], "message": "Привет"}
    ).json()
    response = client.post(
        "/api/chat",
        json={
            "dataset_id": b["dataset_id"],
            "conversation_id": conversation["conversation_id"],
            "message": "Привет",
        },
    )
    assert response.status_code == 409


def test_persistence_after_restart(settings):
    with TestClient(create_app(settings)) as client:
        dataset = upload(client)
        result = calculate(client, dataset["dataset_id"])
    with TestClient(create_app(settings)) as client:
        response = client.get(
            "/api/recommendations", params={"calculation_id": result["calculation_id"]}
        )
        assert response.json()["dataset_id"] == dataset["dataset_id"]
        assert (
            client.get(
                "/api/export", params={"calculation_id": result["calculation_id"]}
            ).status_code
            == 200
        )


def test_provider_receives_exact_parameters(settings):
    adapter = CalculationAdapter(settings)
    original = adapter.module
    received = []

    def calculate_spy(payload, params):
        received.append((payload, params.model_dump(mode="json")))
        return original.calculate(payload, params)

    adapter.module = SimpleNamespace(
        prepare_dataset=original.prepare_dataset, calculate=calculate_spy
    )
    with TestClient(create_app(settings, adapter=adapter)) as client:
        dataset = upload(client)
        wanted = parameters(supplier="ИЭК", demand_growth_adjustment=0.2)
        result = client.post(
            "/api/calculate", json={"dataset_id": dataset["dataset_id"], "parameters": wanted}
        )
        assert result.status_code == 201
        assert received == [({"fixture": "demo-v1"}, wanted)]


def test_provider_mode_change_rejects_old_dataset(settings):
    with TestClient(create_app(settings)) as client:
        dataset = upload(client)
        client.app.state.adapter.name = "new-provider"
        response = client.post(
            "/api/calculate", json={"dataset_id": dataset["dataset_id"], "parameters": parameters()}
        )
        assert response.status_code == 409


def test_invalid_business_data_prevents_calculation(settings):
    adapter = CalculationAdapter(settings)
    original = adapter.module

    def invalid_prepare(paths, sources):
        value = original.prepare_dataset(paths, sources)
        value.report.valid = False
        value.report.errors = ["Не найдена колонка артикула"]
        return value

    adapter.module = SimpleNamespace(prepare_dataset=invalid_prepare, calculate=original.calculate)
    with TestClient(create_app(settings, adapter=adapter)) as client:
        dataset = upload(client)
        assert not dataset["validation"]["valid"]
        response = client.post(
            "/api/calculate", json={"dataset_id": dataset["dataset_id"], "parameters": parameters()}
        )
        assert response.status_code == 422


def test_unknown_ids(client):
    assert client.get("/api/export", params={"calculation_id": str(uuid4())}).status_code == 404
    assert client.get("/api/export", params={"calculation_id": "../../secret"}).status_code == 422


def test_cors(client):
    response = client.options(
        "/api/chat",
        headers={
            "Origin": "http://localhost:5500",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "content-type",
        },
    )
    assert response.headers["access-control-allow-origin"] == "http://localhost:5500"
    response = client.get("/api/health", headers={"Origin": "https://untrusted.example"})
    assert "access-control-allow-origin" not in response.headers


def test_example_contract_valid():
    from pathlib import Path

    from backend.models import Calculation

    value = json.loads(Path("contracts/recommendation.example.json").read_text(encoding="utf-8"))
    assert Calculation.model_validate(value).demo is True
