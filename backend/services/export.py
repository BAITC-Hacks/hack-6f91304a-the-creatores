import json
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "Артикул",
    "Название",
    "Единица",
    "Остаток",
    "Поставки в периоде",
    "Прогноз спроса",
    "Страховой запас",
    "MOQ",
    "Кратность",
    "К заказу",
    "Срочность",
    "Обоснование",
    "Предупреждения",
]
FIELDS = [
    "sku",
    "name",
    "unit",
    "stock",
    "incoming_in_period",
    "forecast_demand",
    "safety_stock",
    "moq",
    "order_multiple",
    "recommended_qty",
    "urgency",
    "reason",
    "warnings",
]


def append_safe(sheet, values):
    sheet.append(
        [
            ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value
            for value in values
        ]
    )
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            # Explicit string cell type prevents =, +, -, @ from becoming formulas.
            cell.data_type = "s"
            cell.number_format = "@"


def format_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245A81")
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(
            65, max(16, width + 2)
        )


def export_calculation(calculation):
    workbook = Workbook()
    parameters = workbook.active
    parameters.title = "Параметры"
    append_safe(parameters, ["Параметр", "Значение"])
    metadata = {
        "calculation_id": str(calculation.calculation_id),
        "dataset_id": str(calculation.dataset_id),
        "Режим": "ДЕМО — синтетические данные" if calculation.demo else "Расчётный модуль",
        "provider": calculation.provider,
        "created_at": calculation.created_at.isoformat(),
        **calculation.parameters.model_dump(mode="json"),
        "forecast_start_inclusive": str(calculation.forecast_start),
        "forecast_end_exclusive": str(calculation.forecast_end),
        "Единицы": "Все количества в единице конкретной строки; null = неизвестно",
        "warnings": "\n".join(calculation.warnings),
    }
    for key, value in metadata.items():
        append_safe(parameters, [key, value])
    for source in calculation.sources:
        append_safe(
            parameters, ["source", json.dumps(source.model_dump(mode="json"), ensure_ascii=False)]
        )
    suppliers = sorted({row.supplier for row in calculation.recommendations})
    for supplier in suppliers:
        title = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "_", supplier).strip("'") or "Поставщик"
        # openpyxl deduplicates titles case-insensitively; reserve room for suffixes.
        sheet = workbook.create_sheet(title[:25])
        append_safe(sheet, HEADERS)
        for row in calculation.recommendations:
            if row.supplier == supplier:
                values = row.model_dump()
                values["warnings"] = "\n".join(row.warnings)
                append_safe(sheet, [values[field] for field in FIELDS])
        append_safe(parameters, [f"Лист: {sheet.title}", supplier])
    for sheet in workbook:
        format_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
