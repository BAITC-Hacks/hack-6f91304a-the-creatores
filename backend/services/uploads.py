import hashlib
import re
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from starlette.concurrency import run_in_threadpool

from backend.errors import AppError
from backend.models import Dataset, Source


def safe_name(name: str | None) -> str:
    value = (name or "").replace("\\", "/").split("/")[-1]
    value = re.sub(r'[\x00-\x1f<>:"|?*]', "_", value).strip(" .")
    if not value or len(value) > 180 or Path(value).suffix.lower() != ".xlsx":
        raise AppError(422, "invalid_file", "Допускаются только файлы .xlsx с корректным именем.")
    return value


def inspect_workbook(path, settings):
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            if len(entries) > 10000 or len(names) != len(entries):
                raise ValueError("Invalid ZIP entries")
            if sum(entry.file_size for entry in entries) > settings.max_uncompressed_bytes:
                raise AppError(413, "workbook_too_large", "Распакованный Excel превышает лимит.")
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise ValueError("Not an XLSX workbook")
            if any("vbaproject" in name.lower() for name in names):
                raise ValueError("Macros are not supported")
            if archive.testzip() is not None:
                raise ValueError("Invalid ZIP checksum")
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if not workbook.sheetnames:
                raise ValueError("Empty workbook")
            cells = 0
            for sheet in workbook:
                # Ignore untrusted dimension metadata; enforce actual parsed cell count.
                sheet.reset_dimensions()
                for row in sheet.iter_rows():
                    cells += len(row)
                    if cells > settings.max_workbook_cells:
                        raise AppError(413, "workbook_too_large", "Слишком много ячеек в Excel.")
            return workbook.sheetnames
        finally:
            workbook.close()
    except AppError:
        raise
    except Exception as exc:
        raise AppError(
            422, "invalid_file", "Файл повреждён или не является допустимым Excel .xlsx."
        ) from exc


async def upload_dataset(files, storage, adapter, settings):
    if not files or len(files) > settings.max_upload_files:
        raise AppError(422, "file_count", f"Загрузите от 1 до {settings.max_upload_files} файлов.")
    dataset_id = uuid4()
    directory = storage.root / "uploads" / str(dataset_id)
    directory.mkdir(parents=True)
    committed = False
    try:
        sources, paths = [], []
        for upload in files:
            filename = safe_name(upload.filename)
            file_id = uuid4()
            path = directory / f"{file_id}.xlsx"
            size, digest = 0, hashlib.sha256()
            with path.open("xb") as target:
                while chunk := await upload.read(64 * 1024):
                    size += len(chunk)
                    if size > settings.max_file_bytes:
                        raise AppError(413, "file_too_large", "Файл превышает допустимый размер.")
                    digest.update(chunk)
                    target.write(chunk)
            sheets = await run_in_threadpool(inspect_workbook, path, settings)
            sources.append(
                Source(
                    file_id=file_id,
                    filename=filename,
                    size_bytes=size,
                    sha256=digest.hexdigest(),
                    sheets=sheets,
                )
            )
            paths.append(path)
        prepared = await run_in_threadpool(adapter.prepare, paths, sources)
        dataset = Dataset(
            dataset_id=dataset_id,
            created_at=datetime.now(timezone.utc),
            demo=adapter.demo,
            sources=sources,
            validation=prepared.report,
            payload=prepared.payload,
            provider=adapter.name,
        )
        storage.create("dataset", dataset_id, dataset)
        committed = True
        return dataset
    finally:
        if not committed:
            shutil.rmtree(directory)
        for upload in files:
            await upload.close()
