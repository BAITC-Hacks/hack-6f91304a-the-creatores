"""Normalize explicit spreadsheet facts; no demo data or silent zero filling.

``aliases`` extends the built-in mapping, e.g. ``{"stock": ["Доступно, шт"]}``.
Supported sheets contain SKU and either product/procurement attributes, long
monthly sales, date-headed wide sales, incoming deliveries, or seasonality.
The first header containing SKU within the first twenty rows is used.
"""

import math
import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from backend.calculations.validation import issue, make_report, validate_dataset

DEFAULT_ALIASES = {
    "sku": ("sku", "артикул", "код", "код товара", "номенклатурный номер", "item code"),
    "name": ("name", "product", "product name", "наименование", "номенклатура", "товар"),
    "supplier": ("supplier", "vendor", "поставщик"),
    "unit": ("unit", "uom", "ед", "ед изм", "единица измерения", "единица"),
    "stock": ("stock", "on hand", "остаток", "остатки", "остаток на складе", "наличие"),
    "moq": ("moq", "минимальный заказ", "минимальная партия", "мин заказ", "минимум заказа"),
    "order_multiple": (
        "order multiple",
        "multiple",
        "pack size",
        "кратность",
        "кратность заказа",
        "упаковка",
    ),
    "lead_time_days": (
        "lead time days",
        "lead time",
        "срок поставки",
        "срок поставки дней",
        "срок доставки",
    ),
    "month": ("month", "period", "sales month", "месяц", "период", "месяц продаж"),
    "date": ("date", "дата"),
    "quantity": ("quantity", "qty", "количество", "кол во", "объем"),
    "sales_quantity": ("sales", "sales quantity", "sales qty", "продажи", "объем продаж"),
    "stockout_days": (
        "stockout days",
        "out of stock days",
        "дни отсутствия",
        "дней без остатка",
        "дни без товара",
        "дней отсутствия товара",
    ),
    "incoming_date": (
        "incoming date",
        "delivery date",
        "arrival date",
        "expected date",
        "eta",
        "дата поставки",
        "дата поступления",
        "ожидаемая дата",
    ),
    "incoming_quantity": (
        "incoming",
        "incoming quantity",
        "incoming qty",
        "in transit",
        "on order",
        "в пути",
        "ожидается",
        "поступление",
        "поступления",
        "количество поставки",
    ),
    "seasonality": (
        "seasonality",
        "seasonality coefficient",
        "seasonal index",
        "coefficient",
        "сезонность",
        "коэффициент сезонности",
        "коэффициент",
        "индекс сезонности",
    ),
}

MONTH_NAMES = (
    ("jan", "january", "янв", "январь", "января"),
    ("feb", "february", "фев", "февраль", "февраля"),
    ("mar", "march", "мар", "март", "марта"),
    ("apr", "april", "апр", "апрель", "апреля"),
    ("may", "май", "мая"),
    ("jun", "june", "июн", "июнь", "июня"),
    ("jul", "july", "июл", "июль", "июля"),
    ("aug", "august", "авг", "август", "августа"),
    ("sep", "sept", "september", "сен", "сент", "сентябрь", "сентября"),
    ("oct", "october", "окт", "октябрь", "октября"),
    ("nov", "november", "ноя", "ноябрь", "ноября"),
    ("dec", "december", "дек", "декабрь", "декабря"),
)


def _normalized(value):
    return re.sub(r"[\W_]+", " ", str(value).casefold().replace("ё", "е")).strip()


def _empty(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_date(value, *, monthly=False):
    if isinstance(value, datetime):
        result = value.date()
    elif isinstance(value, date):
        result = value
    elif isinstance(value, str):
        value = value.strip()
        formats = ["%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"]
        if monthly:
            formats += ["%Y-%m", "%Y/%m", "%m.%Y", "%m/%Y", "%Y.%m"]
        result = None
        for pattern in formats:
            try:
                result = datetime.strptime(value, pattern).date()
                break
            except ValueError:
                pass
        if result is None and monthly:
            parts = _normalized(value).split()
            if len(parts) == 2:
                for number, names in enumerate(MONTH_NAMES, 1):
                    if parts[0] in names and re.fullmatch(r"\d{4}", parts[1]):
                        result = date(int(parts[1]), number, 1)
                    elif parts[1] in names and re.fullmatch(r"\d{4}", parts[0]):
                        result = date(int(parts[0]), number, 1)
        if result is None:
            raise ValueError("Unrecognized date")
    else:
        raise ValueError("Expected an explicit date")
    return result.replace(day=1) if monthly else result


def _month_number(value):
    if isinstance(value, bool):
        raise ValueError("Invalid month")
    if isinstance(value, (int, float)) and math.isfinite(value) and value == int(value):
        if 1 <= value <= 12:
            return str(int(value))
    normalized = _normalized(value)
    if normalized in {str(number) for number in range(1, 13)}:
        return normalized
    for number, names in enumerate(MONTH_NAMES, 1):
        if normalized in names:
            return str(number)
    raise ValueError("Expected month 1..12")


def _number(value):
    if isinstance(value, bool):
        raise ValueError("Boolean is not a quantity")
    if isinstance(value, str):
        value = re.sub(r"[\s\u00a0\u202f]", "", value.strip())
        if "," in value and "." in value:
            raise ValueError("Ambiguous numeric separators")
        value = value.replace(",", ".")
    result = float(value)
    if not math.isfinite(result):
        raise ValueError("Nonfinite number")
    return result


def _sku(cell, value):
    if isinstance(value, bool) or isinstance(value, (datetime, date)):
        raise ValueError("Invalid SKU")
    if isinstance(value, (int, float)):
        if not math.isfinite(value) or value != int(value):
            raise ValueError("Invalid numeric SKU")
        mask = cell.number_format or ""
        if re.fullmatch(r"0+", mask):
            return str(int(value)).zfill(len(mask))
        return str(int(value))
    return str(value).strip()


def load_workbooks(files, sources=None, *, aliases=None):
    """Read Excel files and return ``payload`` plus a structured validation report."""
    files = list(files)
    sources = list(sources or [])
    lookup = {}
    issues = []
    products = {}
    incoming_seen = set()
    sales_seen = set()
    static_seen = set()
    seasonality_seen = set()
    recognized_sheets = 0
    merged_aliases = {key: list(values) for key, values in DEFAULT_ALIASES.items()}
    for field, names in (aliases or {}).items():
        if field not in DEFAULT_ALIASES:
            raise ValueError(f"Unknown canonical alias field: {field}")
        merged_aliases[field].extend([names] if isinstance(names, str) else names)
    for field, names in merged_aliases.items():
        for name in [field, *names]:
            normalized = _normalized(name)
            if normalized in lookup and lookup[normalized] != field:
                raise ValueError(f"Ambiguous column alias: {name}")
            lookup[normalized] = field

    def add(code, severity, message, context, field=None):
        location = f"{context['source']} / {context['sheet']}"
        if context.get("row"):
            location += f" / строка {context['row']}"
        sku = context.get("sku")
        full = f"{location}: {message}"
        issues.append(issue(code, severity, full, **context, field=field))
        if severity == "warning" and sku in products:
            products[sku]["warnings"].append(full)

    def parsed_number(value, context, field):
        if _empty(value):
            return None
        try:
            number = _number(value)
        except (TypeError, ValueError, OverflowError):
            add("invalid_number", "error", f"Некорректное число в поле {field}.", context, field)
            return None
        if number < 0 or (field in ("order_multiple", "seasonality") and number == 0):
            add("invalid_range", "error", f"Недопустимый диапазон поля {field}.", context, field)
            return None
        if field in ("lead_time_days", "stockout_days") and number != int(number):
            add("invalid_integer", "error", f"Поле {field} должно быть целым.", context, field)
            return None
        return number

    def parsed_date(value, context, field, monthly=False):
        if _empty(value):
            return None
        try:
            return _parse_date(value, monthly=monthly).isoformat()
        except (TypeError, ValueError, OverflowError):
            add("invalid_date", "error", f"Некорректная дата в поле {field}.", context, field)
            return None

    def merge(product, field, value, context):
        if value is None:
            return
        if product[field] is not None and product[field] != value:
            add(
                "conflicting_value",
                "error",
                f"Противоречивые значения поля {field} для SKU.",
                context,
                field,
            )
        else:
            product[field] = value

    for file_index, path in enumerate(files):
        metadata = sources[file_index] if file_index < len(sources) else None
        if isinstance(metadata, dict):
            filename = metadata.get("filename") or Path(path).name
        else:
            filename = getattr(metadata, "filename", None) or Path(path).name
        workbook = cached = None
        context = {"source": str(filename), "sheet": None, "row": None, "sku": None}
        try:
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            for sheet in workbook:
                context = {"source": str(filename), "sheet": sheet.title, "row": None, "sku": None}
                sheet.reset_dimensions()
                cached_sheet = cached[sheet.title]
                cached_sheet.reset_dimensions()
                role_hint = _normalized(f"{filename} {sheet.title}")
                incoming_role = any(
                    word in role_hint
                    for word in (
                        "incoming",
                        "deliver",
                        "arrival",
                        "transit",
                        "поступ",
                        "в пути",
                        "поставки",
                    )
                )
                seasonal_role = "season" in role_hint or "сезон" in role_hint
                stock_role = any(word in role_hint for word in ("stock", "остат", "склад"))
                columns = None
                wide_months = {}
                seasonal_months = {}
                wide_stockouts = {}
                had_content = False
                for row_number, (row, cached_row) in enumerate(
                    zip(sheet.iter_rows(), cached_sheet.iter_rows(), strict=True), 1
                ):
                    if not any(not _empty(cell.value) for cell in row):
                        continue
                    had_content = True
                    if columns is None:
                        if row_number > 20:
                            break
                        candidates = {}
                        for column, cell in enumerate(row):
                            field = (
                                lookup.get(_normalized(cell.value))
                                if cell.value is not None
                                else None
                            )
                            if field:
                                candidates.setdefault(field, []).append(column)
                        if "sku" not in candidates:
                            continue
                        columns = {field: indices[0] for field, indices in candidates.items()}
                        for field, indices in candidates.items():
                            if len(indices) > 1:
                                add(
                                    "duplicate_column",
                                    "error",
                                    f"Поле {field} имеет несколько колонок.",
                                    {**context, "row": row_number},
                                    field,
                                )
                        sheet_incoming = incoming_role or bool(
                            set(columns) & {"incoming_date", "incoming_quantity"}
                        )

                        def remap(field, target):
                            if target in columns:
                                add(
                                    "duplicate_column",
                                    "error",
                                    f"Колонки {field} и {target} обозначают одно поле {target}.",
                                    {**context, "row": row_number},
                                    target,
                                )
                            else:
                                columns[target] = columns.pop(field)

                        if "date" in columns:
                            target = "incoming_date" if sheet_incoming else "month"
                            remap("date", target)
                        if "quantity" in columns:
                            target = (
                                "incoming_quantity"
                                if sheet_incoming
                                else "stock"
                                if stock_role
                                else "sales_quantity"
                            )
                            remap("quantity", target)
                        for column, cell in enumerate(row):
                            if column in columns.values() or _empty(cell.value):
                                continue
                            if seasonal_role:
                                try:
                                    seasonal_months[column] = _month_number(cell.value)
                                    continue
                                except ValueError:
                                    pass
                            try:
                                wide_months[column] = _parse_date(
                                    cell.value, monthly=True
                                ).isoformat()
                                continue
                            except (TypeError, ValueError, OverflowError):
                                pass
                            normalized = _normalized(cell.value)
                            for prefix in merged_aliases["stockout_days"]:
                                prefix = _normalized(prefix)
                                if normalized.startswith(prefix + " "):
                                    # Keep punctuation in the date after the alias.
                                    suffix = re.sub(r"^[^0-9]*", "", str(cell.value)).strip()
                                    try:
                                        wide_stockouts[column] = _parse_date(
                                            suffix, monthly=True
                                        ).isoformat()
                                    except (TypeError, ValueError, OverflowError):
                                        pass
                        if not (
                            set(columns) - {"sku", "name", "supplier", "unit"}
                            or wide_months
                            or seasonal_months
                        ):
                            add(
                                "metadata_only_sheet",
                                "warning",
                                "Лист содержит только карточки товаров.",
                                {**context, "row": row_number},
                            )
                        recognized_sheets += 1
                        continue

                    context = {**context, "row": row_number, "sku": None}
                    values = []
                    for column, cell in enumerate(row):
                        value = cell.value
                        if cell.data_type == "f":
                            value = cached_row[column].value
                            if value is None:
                                field = next((f for f, c in columns.items() if c == column), None)
                                add(
                                    "formula_without_cache",
                                    "warning",
                                    "У формулы нет сохранённого результата; значение неизвестно.",
                                    context,
                                    field,
                                )
                        elif cell.data_type == "e":
                            field = next((f for f, c in columns.items() if c == column), None)
                            add(
                                "excel_cell_error",
                                "error",
                                "Ячейка содержит ошибку Excel.",
                                context,
                                field,
                            )
                            value = None
                        values.append(value)

                    def get(field):
                        column = columns.get(field)
                        return (
                            values[column] if column is not None and column < len(values) else None
                        )

                    raw_sku = get("sku")
                    if _empty(raw_sku):
                        add(
                            "missing_sku",
                            "error",
                            "В непустой строке отсутствует SKU.",
                            context,
                            "sku",
                        )
                        continue
                    try:
                        sku = _sku(row[columns["sku"]], raw_sku)
                    except (ValueError, OverflowError):
                        add("invalid_sku", "error", "Некорректный SKU.", context, "sku")
                        continue
                    context["sku"] = sku
                    product = products.setdefault(
                        sku,
                        {
                            "sku": sku,
                            "name": None,
                            "supplier": None,
                            "unit": None,
                            "stock": None,
                            "moq": None,
                            "order_multiple": None,
                            "lead_time_days": None,
                            "sales": [],
                            "incoming": [],
                            "incoming_known": False,
                            "seasonality": {},
                            "warnings": [],
                        },
                    )
                    for field in ("name", "supplier", "unit"):
                        value = get(field)
                        merge(
                            product, field, None if _empty(value) else str(value).strip(), context
                        )
                    for field in ("stock", "moq", "order_multiple", "lead_time_days"):
                        merge(product, field, parsed_number(get(field), context, field), context)
                    is_timeseries = bool(
                        wide_months
                        or seasonal_months
                        or set(columns)
                        & {
                            "incoming_date",
                            "incoming_quantity",
                            "sales_quantity",
                            "seasonality",
                        }
                    )
                    if not is_timeseries:
                        key = (file_index, sheet.title, sku)
                        if key in static_seen:
                            add(
                                "duplicate_sku",
                                "error",
                                "SKU повторяется в листе карточек товаров, остатков или закупок.",
                                context,
                                "sku",
                            )
                        static_seen.add(key)

                    def sale(month, quantity, stockout):
                        if month is None:
                            add(
                                "missing_sales_month",
                                "error",
                                "Не указан месяц продаж.",
                                context,
                                "month",
                            )
                            return
                        key = (sku, month)
                        if key in sales_seen:
                            add(
                                "duplicate_sales",
                                "error",
                                "Продажи SKU за месяц указаны повторно.",
                                context,
                                "month",
                            )
                            return
                        sales_seen.add(key)
                        product["sales"].append(
                            {
                                "month": month,
                                "quantity": parsed_number(quantity, context, "sales.quantity"),
                                "stockout_days": parsed_number(stockout, context, "stockout_days"),
                            }
                        )

                    if "sales_quantity" in columns:
                        sale(
                            parsed_date(get("month"), context, "month", monthly=True),
                            get("sales_quantity"),
                            get("stockout_days"),
                        )
                    for column, month in wide_months.items():
                        stockout_column = next(
                            (c for c, m in wide_stockouts.items() if m == month), None
                        )
                        stockout = (
                            values[stockout_column]
                            if stockout_column is not None and stockout_column < len(values)
                            else None
                        )
                        sale(month, values[column] if column < len(values) else None, stockout)
                    if "incoming_quantity" in columns or "incoming_date" in columns:
                        incoming_seen.add(sku)
                        product["incoming"].append(
                            {
                                "date": parsed_date(get("incoming_date"), context, "incoming.date"),
                                "quantity": parsed_number(
                                    get("incoming_quantity"), context, "incoming.quantity"
                                ),
                            }
                        )

                    def seasonal(month, value):
                        key = (sku, month)
                        if key in seasonality_seen:
                            add(
                                "duplicate_seasonality",
                                "error",
                                "Коэффициент месяца указан повторно.",
                                context,
                                "seasonality",
                            )
                            return
                        seasonality_seen.add(key)
                        coefficient = parsed_number(value, context, "seasonality")
                        if coefficient is not None:
                            product["seasonality"][month] = coefficient

                    if "seasonality" in columns:
                        try:
                            seasonal(_month_number(get("month")), get("seasonality"))
                        except ValueError:
                            add(
                                "invalid_seasonality_month",
                                "error",
                                "Месяц сезонности должен быть 1..12.",
                                context,
                                "month",
                            )
                    for column, month in seasonal_months.items():
                        seasonal(month, values[column] if column < len(values) else None)
                if columns is None:
                    add(
                        "missing_sku_header" if had_content else "empty_sheet",
                        "error" if had_content else "warning",
                        "Не найдена колонка SKU в первых 20 строках."
                        if had_content
                        else "Пустой лист.",
                        context,
                        "sku",
                    )
        except Exception as exc:
            # Parsing failures remain reviewable business reports; do not leak workbook contents.
            add(
                "workbook_read_error",
                "error",
                f"Не удалось прочитать Excel ({type(exc).__name__}).",
                context,
            )
        finally:
            if workbook is not None:
                workbook.close()
            if cached is not None:
                cached.close()

    for sku, product in products.items():
        product["sales"].sort(key=lambda record: record["month"])
        product["incoming_known"] = sku in incoming_seen and all(
            item["quantity"] is not None and (item["quantity"] == 0 or item["date"] is not None)
            for item in product["incoming"]
        )
    payload = {
        "schema_version": 1,
        "products": list(products.values()),
        "validation_issues": list(issues),
    }
    validation = validate_dataset(payload)
    issues = validation["summary"]["issues"]
    for item in validation["summary"]["issues"]:
        if item["severity"] == "warning" and item["sku"] in products:
            products[item["sku"]]["warnings"].append(item["message"])
    return {
        "payload": payload,
        "report": make_report(
            issues,
            products=len(products),
            uploaded_files=len(files),
            recognized_sheets=recognized_sheets,
        ),
    }
